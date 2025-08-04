#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
測試Word文檔分割腳本
創建一個測試用的Word文檔，然後測試分割功能
"""

import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_test_docx():
    """
    創建一個測試用的Word文檔，模擬BallonTranslator的輸出格式
    """
    doc = Document()
    
    # 設置文檔樣式
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = 12
    
    # 頁面1
    doc.add_paragraph("page_001.jpg")
    table1 = doc.add_table(rows=2, cols=2, style='Table Grid')
    table1.cell(0, 0).text = "氣泡1圖片"
    table1.cell(0, 1).text = "這是第一頁的第一個氣泡翻譯"
    table1.cell(1, 0).text = "氣泡2圖片"
    table1.cell(1, 1).text = "這是第一頁的第二個氣泡翻譯"
    doc.add_page_break()
    
    # 頁面2
    doc.add_paragraph("page_002.jpg")
    table2 = doc.add_table(rows=1, cols=2, style='Table Grid')
    table2.cell(0, 0).text = "氣泡3圖片"
    table2.cell(0, 1).text = "這是第二頁的氣泡翻譯"
    doc.add_page_break()
    
    # 頁面3
    doc.add_paragraph("page_003.jpg")
    table3 = doc.add_table(rows=3, cols=2, style='Table Grid')
    table3.cell(0, 0).text = "氣泡4圖片"
    table3.cell(0, 1).text = "這是第三頁的第一個氣泡翻譯"
    table3.cell(1, 0).text = "氣泡5圖片"
    table3.cell(1, 1).text = "這是第三頁的第二個氣泡翻譯"
    table3.cell(2, 0).text = "氣泡6圖片"
    table3.cell(2, 1).text = "這是第三頁的第三個氣泡翻譯"
    
    return doc

def test_split_functionality():
    """
    測試分割功能
    """
    print("🧪 開始測試Word文檔分割功能...")
    
    # 創建測試文檔
    test_doc = create_test_docx()
    test_file = "test_ballontranslator.docx"
    test_doc.save(test_file)
    print(f"✅ 已創建測試文檔: {test_file}")
    
    # 測試分割功能
    try:
        # 導入分割函數
        from split_ballontranslator_docx import split_ballontranslator_docx
        
        # 執行分割
        success = split_ballontranslator_docx(test_file, "test_output")
        
        if success:
            print("✅ 分割測試成功！")
            
            # 檢查輸出文件
            output_dir = Path("test_output")
            if output_dir.exists():
                files = list(output_dir.glob("*.docx"))
                print(f"📁 生成的文件數量: {len(files)}")
                for file in files:
                    print(f"  📄 {file.name}")
                
                # 檢查摘要文件
                summary_file = output_dir / "分割摘要.txt"
                if summary_file.exists():
                    print(f"📋 摘要文件已生成: {summary_file}")
                else:
                    print("⚠️  摘要文件未生成")
            else:
                print("❌ 輸出目錄未創建")
        else:
            print("❌ 分割測試失敗！")
            
    except ImportError as e:
        print(f"❌ 無法導入分割模塊: {e}")
        print("請確保 split_ballontranslator_docx.py 文件在同一目錄下")
    except Exception as e:
        print(f"❌ 測試過程中發生錯誤: {e}")
        import traceback
        traceback.print_exc()
    
    # 清理測試文件
    try:
        if os.path.exists(test_file):
            os.remove(test_file)
            print(f"🗑️  已清理測試文件: {test_file}")
    except Exception as e:
        print(f"⚠️  清理測試文件時發生錯誤: {e}")

def main():
    """主函數"""
    print("=" * 60)
    print("Word文檔分割腳本測試工具")
    print("=" * 60)
    
    # 檢查依賴
    try:
        import docx
        print("✅ python-docx 庫已安裝")
    except ImportError:
        print("❌ 未安裝 python-docx 庫")
        print("請運行: pip install python-docx")
        return
    
    # 執行測試
    test_split_functionality()
    
    print("\n" + "=" * 60)
    print("測試完成！")
    print("=" * 60)

if __name__ == "__main__":
    main() 